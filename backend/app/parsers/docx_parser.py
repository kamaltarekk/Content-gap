from __future__ import annotations

import io

from app.parsers.base import ParsedDocument


def parse_docx(data: bytes, filename: str | None = None) -> ParsedDocument:
    from docx import Document

    try:
        doc = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(
            text="",
            extraction_quality="unreadable",
            warnings=[f"docx_open_failed:{type(exc).__name__}"],
            parser_name="python-docx",
            parser_version="1",
            title=filename,
        )
    paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
    text = "\n\n".join(paragraphs)
    title = paragraphs[0] if paragraphs else filename
    quality = "complete" if text.strip() else "unreadable"
    return ParsedDocument(
        text=text,
        title=title,
        extraction_quality=quality,
        parser_name="python-docx",
        parser_version="1",
        metadata={"paragraph_count": len(paragraphs)},
    )
